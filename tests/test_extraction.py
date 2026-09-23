import asyncio
import base64
import json
from dataclasses import replace
from io import BytesIO
from unittest.mock import AsyncMock
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
from docx import Document
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from botkit.extraction import (
    ContentExtractionService,
    ExtractedContent,
    ExtractionConfig,
    ExtractionError,
    ExtractionLimitError,
    ImageExtractor,
    UnsupportedContentError,
    render_pdf,
)
from botkit.extraction.documents import DOCX, PPTX, XLSX
from botkit.llm import FallbackLLM, LLMConfig, LLMGateway, LLMResponse
from botkit.runtime import ChatBot
from botkit.transport import Attachment, IncomingMessage, MemoryAdapter
from botkit.usage import SQLiteUsageTracker, usage_context


def png(size=(120, 60)):
    buffer = BytesIO()
    with Image.new("RGB", size, "white") as image:
        image.save(buffer, format="PNG")
    return buffer.getvalue()


def pdf(kinds=("text",)):
    buffer = BytesIO()
    doc = canvas.Canvas(buffer, pagesize=(400, 500))
    for index, kind in enumerate(kinds, 1):
        if kind in {"text", "mixed"}:
            doc.drawString(
                30, 450, f"Page {index}: project owners and responsibilities in the supplied document."
            )
        if kind in {"scan", "mixed"}:
            doc.drawImage(ImageReader(BytesIO(png())), 20, 40, width=360, height=350)
        doc.showPage()
    doc.save()
    return buffer.getvalue()


class Vision:
    model_name = "test-vision"

    def __init__(self):
        self.calls = []

    async def ainvoke_with_image_b64(self, prompt, image_b64, **options):
        assert base64.b64decode(image_b64).startswith(b"\x89PNG")
        assert "stream" not in options
        self.calls.append((prompt, options))
        payload = {
            "text": "Recognized page",
            "tables": [{"rows": [["A", "B"], ["1", "2"]]}],
            "warnings": ["uncertain_cell"],
        }
        result = LLMResponse(
            json.dumps(payload), json.dumps(payload), "raw_text_fallback", model=self.model_name
        )
        options["validator"](result)
        return result


async def test_text_pdf_never_calls_llm():
    vision = Vision()
    result = await ContentExtractionService(vision).extract_bytes(pdf(), "application/pdf")
    assert "project owners" in result.text
    assert result.metadata["vision_pages"] == 0
    assert result.source_confidence == "high" and not vision.calls


async def test_mixed_pdf_routes_per_page_preserves_order_and_warnings():
    vision = Vision()
    result = await ContentExtractionService(vision).extract_bytes(pdf(("text", "scan")), "application/pdf")
    assert len(vision.calls) == 1
    assert [p["method"] for p in result.structured_data["pages"]] == ["text", "vision"]
    assert result.text.index("project owners") < result.text.index("Recognized page")
    assert result.source_confidence == "low"
    assert "page:2:uncertain_cell" in result.warnings
    assert result.structured_data["pages"][1]["tables"][0]["rows"][1] == ["1", "2"]


async def test_text_pdf_with_large_image_also_uses_vision_in_auto():
    vision = Vision()
    result = await ContentExtractionService(vision).extract_bytes(pdf(("mixed",)), "application/pdf")
    assert len(vision.calls) == 1 and result.metadata["vision_pages"] == 1


@pytest.mark.parametrize("mode,expected", [("text", 0), ("vision", 2)])
async def test_explicit_pdf_modes(mode, expected):
    vision = Vision()
    result = await ContentExtractionService(vision, config=ExtractionConfig(pdf_mode=mode)).extract_bytes(
        pdf(("text", "scan")), "application/pdf"
    )
    assert len(vision.calls) == expected
    if mode == "text":
        assert "page:2:no_text_layer" in result.warnings


async def test_missing_vision_reports_empty_scanned_page_not_success_silently():
    result = await ContentExtractionService().extract_bytes(pdf(("scan",)), "application/pdf")
    assert "page:1:vision_unavailable" in result.warnings
    with pytest.raises(ExtractionError, match="requires"):
        await ContentExtractionService(config=ExtractionConfig(pdf_mode="vision")).extract_bytes(
            pdf(), "application/pdf"
        )


async def test_pdf_limits_fail_before_paid_calls():
    vision = Vision()
    for config in [ExtractionConfig(max_pages=1), ExtractionConfig(max_vision_pages=1)]:
        with pytest.raises(ExtractionLimitError):
            await ContentExtractionService(vision, config=config).extract_bytes(
                pdf(("scan", "scan")), "application/pdf"
            )
    assert not vision.calls


async def test_renderer_returns_actual_bounded_pngs():
    config = ExtractionConfig(render_max_pixels=50_000)
    pages = await render_pdf(pdf(("text", "scan")), config=config)
    assert [p.page_number for p in pages] == [1, 2]
    for page in pages:
        with Image.open(BytesIO(page.data)) as image:
            assert image.format == "PNG" and image.size == (page.width, page.height)
            assert image.width * image.height <= config.render_max_pixels
    with pytest.raises(ExtractionLimitError):
        await render_pdf(pdf(), config=replace(config, max_rendered_bytes=10))


async def test_concurrent_pdf_rendering_is_safe():
    outputs = await asyncio.gather(*(render_pdf(pdf()) for _ in range(4)))
    assert all(len(pages) == 1 for pages in outputs)


async def test_vision_and_rendering_do_not_depend_on_native_text_parser(monkeypatch):
    import botkit.extraction.pdf as pdf_module

    def broken(*args):
        raise ExtractionError("Native parser failure")

    monkeypatch.setattr(pdf_module, "_native_pages", broken)
    assert len(await render_pdf(pdf())) == 1
    for mode in ("auto", "vision"):
        result = await ContentExtractionService(
            Vision(), config=ExtractionConfig(pdf_mode=mode)
        ).extract_bytes(pdf(), "application/pdf")
        assert result.metadata["vision_pages"] == 1
        if mode == "auto":
            assert "page:1:native_parser_failed" in result.warnings


async def test_pdf_native_tables_share_the_vision_rows_contract():
    buffer = BytesIO()
    document = canvas.Canvas(buffer, pagesize=(300, 300))
    for x in (20, 140, 260):
        document.line(x, 180, x, 260)
    for y in (180, 220, 260):
        document.line(20, y, 260, y)
    for x, y, text in [(30, 240, "Name"), (150, 240, "Role"), (30, 200, "Alice"), (150, 200, "Owner")]:
        document.drawString(x, y, text)
    document.showPage()
    document.save()
    result = await ContentExtractionService(config=ExtractionConfig(pdf_mode="text")).extract_bytes(
        buffer.getvalue(), "application/pdf"
    )
    assert result.structured_data["pages"][0]["tables"][0]["rows"] == [["Name", "Role"], ["Alice", "Owner"]]


@pytest.mark.parametrize("data", [b"%PDF-1.7 broken", b"not a pdf"])
async def test_bad_pdf_is_an_explicit_error(data):
    with pytest.raises(ExtractionError):
        await ContentExtractionService().extract_bytes(data, "application/pdf")


async def test_password_pdf_is_an_explicit_error():
    buffer = BytesIO()
    document = canvas.Canvas(buffer, encrypt="secret-password")
    document.drawString(10, 10, "Secret document")
    document.save()
    with pytest.raises(ExtractionError, match="encrypted"):
        await ContentExtractionService().extract_bytes(buffer.getvalue(), "application/pdf")


async def test_image_sniffing_size_limits_and_low_confidence():
    vision = Vision()
    result = await ContentExtractionService(vision).extract_bytes(
        png(), "application/octet-stream", filename="upload.bin"
    )
    assert result.kind == "structured" and result.source_confidence == "low"
    assert "mime_type_inferred_from_content_or_filename" in result.warnings
    with pytest.raises(ExtractionLimitError):
        await ImageExtractor(vision, ExtractionConfig(max_image_pixels=10)).extract(png(), "image/png")
    with pytest.raises(ExtractionError):
        await ImageExtractor(vision).extract(b"not an image", "image/png")


def test_transparent_image_is_composited_on_white():
    from botkit.extraction.vision import normalize_image

    buffer = BytesIO()
    with Image.new("RGBA", (20, 10), (0, 0, 0, 0)) as original:
        original.save(buffer, format="PNG")
    image, _ = normalize_image(buffer.getvalue(), ExtractionConfig())
    with Image.open(BytesIO(image)) as decoded:
        assert decoded.getpixel((0, 0)) == (255, 255, 255)


async def test_vision_schema_failure_uses_a1_fallback_and_records_both_calls(tmp_path):
    tracker = SQLiteUsageTracker(tmp_path / "usage.sqlite3")
    calls = []

    def endpoint(request):
        body = json.loads(request.content)
        calls.append(body["model"])
        assert body["stream"] is False
        assert body["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,")
        text = (
            "invalid json"
            if body["model"] == "qwen"
            else json.dumps({"text": "OCR", "tables": [], "warnings": []})
        )
        return httpx.Response(
            200,
            json={
                "model": body["model"],
                "choices": [{"message": {"content": text}}],
                "usage": {"prompt_tokens": 10, "completion_tokens": 5},
            },
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(endpoint)) as client:
        clients = [
            LLMGateway(LLMConfig("hub", name, "https://vision.test/v1"), client=client, tracker=tracker)
            for name in ["qwen", "backup"]
        ]
        chain = FallbackLLM(clients[0], clients[1:], tracker=tracker)
        with usage_context(user_id="u", skill_context="OCR"):
            result = await ContentExtractionService(chain, tracker=tracker).extract_bytes(png(), "image/png")
    assert result.text == "OCR" and calls == ["qwen", "backup"]
    assert result.metadata["fallback_used"] is True
    assert (await tracker.stats()).calls == 2
    events = await tracker.events(kind="extraction")
    assert events[0]["user_id"] == "u" and events[0]["meta"]["warning_count"] == 1
    assert "OCR" not in events[0]["output_text"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32", "cp1251"])
async def test_text_encodings(encoding):
    result = await ContentExtractionService().extract_bytes("Привет, таблица".encode(encoding), "text/plain")
    assert result.text == "Привет, таблица"
    assert bool(result.warnings) == (encoding == "cp1251")


async def test_json_html_and_csv_without_network():
    service = ContentExtractionService()
    result = await service.extract_bytes(b'{"ok":true}', "application/json")
    assert result.structured_data == {"value": {"ok": True}}
    broken = await service.extract_bytes(b'{"broken', "application/json")
    assert "invalid_json_preserved_as_text" in broken.warnings
    html = await service.extract_bytes(b"<script>bad()</script><p>Hello &amp; world</p>", "text/html")
    assert "bad" not in html.text and "Hello & world" in html.text
    csv = await service.extract_bytes(b"a,b\n1,2", "application/octet-stream", filename="table.csv")
    assert "1,2" in csv.text


async def test_docx_keeps_paragraph_table_order_and_structured_cells():
    doc = Document()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Role A", "Role B"
    doc.add_paragraph("After table")
    buffer = BytesIO()
    doc.save(buffer)
    result = await ContentExtractionService().extract_bytes(buffer.getvalue(), "application/octet-stream")
    assert result.text.index("Before") < result.text.index("Role A") < result.text.index("After")
    assert result.structured_data["tables"][0]["rows"] == [["Role A", "Role B"]]
    assert result.metadata["mime_type"] == DOCX


async def test_xlsx_formula_is_not_executed_and_sheet_names_are_kept():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Name", "Value"])
    sheet.append(["a", "=1+2"])
    buffer = BytesIO()
    workbook.save(buffer)
    result = await ContentExtractionService().extract_bytes(buffer.getvalue(), XLSX)
    assert result.structured_data["sheets"][0]["name"] == "Budget"
    assert "=1+2" in result.text and "xlsx_formulas_not_evaluated" in result.warnings


async def test_pptx_text_and_slide_numbers():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Project status"
    buffer = BytesIO()
    presentation.save(buffer)
    result = await ContentExtractionService().extract_bytes(buffer.getvalue(), PPTX)
    assert "Project status" in result.text
    assert result.structured_data["slides"][0]["slide_number"] == 1


@pytest.mark.parametrize(
    "content,config",
    [
        (b"A" * 200, ExtractionConfig(max_archive_bytes=100)),
        (b'<!DOCTYPE x [<!ENTITY s SYSTEM "file:///secret">]><x>&s;</x>', ExtractionConfig()),
    ],
)
async def test_unsafe_ooxml_is_rejected(content, config):
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", content)
    with pytest.raises(ExtractionError):
        await ContentExtractionService(config=config).extract_bytes(buffer.getvalue(), DOCX)


async def test_attachment_limits_checked_before_download():
    reader = AsyncMock(return_value=b"too big")
    service = ContentExtractionService(config=ExtractionConfig(max_bytes=2))
    with pytest.raises(ExtractionLimitError):
        await service.extract(Attachment("document", "text/plain", reader, size=3))
    reader.assert_not_called()
    with pytest.raises(ExtractionLimitError):
        await service.extract(Attachment("document", "text/plain", reader))


async def test_unknown_formats_and_binary_text_are_not_silently_accepted():
    for data, mime in [(b"binary", "application/x-unknown"), (b"\xd0\xcf\x11\xe0rest", "application/msword")]:
        with pytest.raises(UnsupportedContentError):
            await ContentExtractionService().extract_bytes(data, mime)
    with pytest.raises(ExtractionError):
        await ContentExtractionService().extract_bytes(b"a\x00b", "text/plain")


async def test_custom_extractor_takes_precedence_and_retains_warnings():
    extractor = AsyncMock()
    extractor.supports.return_value = True
    extractor.extract.return_value = ExtractedContent("text", "custom", None, ["custom_warning"], "unknown")
    service = ContentExtractionService()
    service.register(extractor)
    result = await service.extract_bytes(b"hi", "text/plain")
    assert result.text == "custom" and result.warnings == ["custom_warning"]


async def test_extraction_cancellation_propagates():
    provider = Vision()
    provider.ainvoke_with_image_b64 = AsyncMock(side_effect=asyncio.CancelledError)
    with pytest.raises(asyncio.CancelledError):
        await ContentExtractionService(provider).extract_bytes(png(), "image/png")


async def test_attachment_only_message_reaches_chat_with_extracted_content():
    llm = AsyncMock()
    llm.achat.return_value = LLMResponse("Reviewed", "Reviewed", "raw_text_fallback")
    adapter = MemoryAdapter()
    ChatBot(llm, [adapter], extraction=ContentExtractionService())
    await adapter.dispatch(
        IncomingMessage(
            "memory",
            "u",
            "c",
            attachments=[Attachment("document", "text/plain", AsyncMock(return_value=b"Document facts"))],
        )
    )
    assert adapter.sent[-1]["text"] == "Reviewed"
    messages = llm.achat.call_args.args[0]
    assert "Document facts" in messages[1]["content"]
    assert "недоверенные" in messages[0]["content"]
    await adapter.aclose()


async def test_audio_uses_optional_a1_transcriber_without_inventing_vision_support(tmp_path):
    tracker = SQLiteUsageTracker(tmp_path / "audio.sqlite3")
    audio = b"RIFF-audio-test-data"

    def endpoint(request):
        assert request.url.path == "/v1/audio/transcriptions"
        assert request.headers["content-type"].startswith("multipart/form-data")
        assert audio in request.content and b'filename="audio.wav"' in request.content
        return httpx.Response(200, json={"text": "Recognized speech"})

    async with httpx.AsyncClient(transport=httpx.MockTransport(endpoint)) as client:
        asr = LLMGateway(
            LLMConfig("openai_compatible", "asr-model", "https://asr.test/v1"), tracker=tracker, client=client
        )
        result = await ContentExtractionService(transcriber=asr).extract_bytes(audio, "audio/wav")
    assert result.kind == "transcript" and result.text == "Recognized speech"
    assert result.source_confidence == "low"
    (event,) = await tracker.events()
    assert event["operation"] == "transcription" and event["usage_known"] is False
    assert "RIFF" not in json.dumps(event)
    with pytest.raises(UnsupportedContentError, match="ASR"):
        await ContentExtractionService(Vision()).extract_bytes(audio, "audio/wav")
