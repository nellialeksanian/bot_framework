"""Operator-run acceptance checks. No network without --live; never authorizes a merge."""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import json
import logging
import os
import sys
import time
from dataclasses import asdict, replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import httpx
from dotenv import load_dotenv

from botkit.extraction import ContentExtractionService, ExtractionConfig, render_pdf
from botkit.extraction.documents import DOCX
from botkit.llm import FallbackLLM, LLMGateway, load_llm
from botkit.transport import Button
from botkit.usage import PriceBook, SQLiteUsageTracker, usage_context

ROOT = Path(__file__).resolve().parents[1]
MARKER = "BOTKIT-7319"
KNOWN_LIMITS = [
    "This is smoke acceptance, not full domain or all-method API certification.",
    "Native-handler interception and VK callback message IDs need regression checks.",
    "Messenger attachment normalization does not cover all platform media types.",
    "Cost must be reconciled with provider billing; unknown cost is not zero.",
]


class CheckFailure(Exception):
    """Only fixed, developer-controlled codes may be passed here (not API responses)."""


def require(condition, code):
    if not condition:
        raise CheckFailure(code)


class Report:
    def __init__(self, directory: Path, mode: str):
        self.directory = directory
        directory.mkdir(parents=True, exist_ok=False)
        self.data = {
            "started_at": datetime.now(UTC).isoformat(),
            "mode": mode,
            "cases": [],
            "required": [],
            "publication_allowed": False,
            "known_limitations": KNOWN_LIMITS,
        }
        self.save()

    def expect(self, *names):
        self.data["required"] = sorted(set(self.data["required"]) | set(names))
        self.save()

    def save(self):
        passed = {c["name"] for c in self.data["cases"] if c["status"] == "passed"}
        self.data["missing"] = sorted(set(self.data["required"]) - passed)
        failed = any(c["status"] == "failed" for c in self.data["cases"])
        self.data["result"] = "failed" if failed else "incomplete" if self.data["missing"] else "passed"
        target = self.directory / "report.json"
        temporary = target.with_suffix(".tmp")
        temporary.write_text(
            json.dumps(self.data, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )
        temporary.replace(target)

    def record(self, name, status="passed", **details):
        self.data["cases"].append({"name": name, "status": status, **details})
        self.save()
        reason = details.get("reason") or details.get("error_type")
        suffix = f" ({reason})" if status == "failed" and reason else ""
        print(f"[{status.upper()}] {name}{suffix}", flush=True)

    async def check(self, name, operation, *, timeout=240):
        started = time.monotonic()
        try:
            with usage_context(bot_id="acceptance", skill_context=name):
                async with asyncio.timeout(timeout):
                    details = await operation()
            self.record(name, elapsed_seconds=round(time.monotonic() - started, 3), **(details or {}))
            return True
        except Exception as exc:
            # Exception messages/tracebacks can contain tokens, signed URLs or document text.
            details = {"error_type": type(exc).__name__}
            causes, cause = [], exc
            while cause is not None and len(causes) < 8:
                causes.append(type(cause).__name__)
                cause = cause.__cause__ or cause.__context__
            details["cause_types"] = causes
            if isinstance(exc, CheckFailure):
                details["reason"] = str(exc)
            status = getattr(exc, "status_code", None)
            if isinstance(status, int):
                details["http_status"] = status
            self.record(name, "failed", **details)
            return False

    @property
    def exit_code(self):
        return 0 if self.data["result"] == "passed" else 1


def make_fixtures(directory: Path):
    """Small synthetic fixtures only; no user documents are uploaded automatically."""
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    directory.mkdir()
    image = Image.new("RGB", (900, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 40), MARKER, fill="black", font=ImageFont.load_default(size=64))
    draw.text((40, 150), "TOTAL 42", fill="black", font=ImageFont.load_default(size=48))
    png = BytesIO()
    image.save(png, format="PNG")
    image.close()

    def pdf(scan):
        output = BytesIO()
        doc = canvas.Canvas(output, pagesize=(600, 400))
        if scan:
            doc.drawImage(ImageReader(BytesIO(png.getvalue())), 30, 150, width=540, height=180)
        else:
            doc.setFont("Helvetica", 20)
            doc.drawString(30, 320, f"{MARKER} - synthetic acceptance document, TOTAL 42.")
        doc.showPage()
        doc.save()
        return output.getvalue()

    docx = BytesIO()
    doc = Document()
    doc.add_paragraph(f"{MARKER} synthetic acceptance document")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "TOTAL", "42"
    doc.save(docx)
    fixtures = {
        "sample.txt": (f"{MARKER} TOTAL 42".encode(), "text/plain"),
        "sample.docx": (docx.getvalue(), DOCX),
        "text.pdf": (pdf(False), "application/pdf"),
        "scan.pdf": (pdf(True), "application/pdf"),
        "sample.png": (png.getvalue(), "image/png"),
    }
    for name, (data, _) in fixtures.items():
        (directory / name).write_bytes(data)
    return fixtures


def content_details(result):
    return {
        "text_chars": len(result.text or ""),
        "warning_count": len(result.warnings),
        "confidence": result.source_confidence,
        "vision_pages": result.metadata.get("vision_pages"),
    }


async def local_checks(report, fixtures, tracker):
    service = ContentExtractionService(config=ExtractionConfig(pdf_mode="text"), tracker=tracker)
    for name in ("sample.txt", "sample.docx", "text.pdf"):
        case = "local:" + name
        report.expect(case)

        async def extract(name=name):
            data, mime = fixtures[name]
            result = await service.extract_bytes(data, mime)
            require(MARKER in (result.text or ""), "fixture_marker_not_extracted")
            return content_details(result)

        await report.check(case, extract)

    async def render():
        pages = await render_pdf(fixtures["scan.pdf"][0])
        require(len(pages) == 1 and pages[0].data.startswith(b"\x89PNG"), "pdf_render_failed")
        (report.directory / "fixtures" / "rendered-scan.png").write_bytes(pages[0].data)
        return {"pages": 1}

    report.expect("local:pdf_to_png")
    await report.check("local:pdf_to_png", render)


async def regression_checks(report):
    """Exercise real SDK dispatch/serialization without contacting either messenger."""
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from aiogram import Bot
    from aiogram.types import Message, Update
    from vkbottle import API
    from vkbottle import Bot as VKBot

    from botkit.transport.telegram import TelegramAdapter
    from botkit.transport.vk import VKAdapter

    payload = {
        "message_id": 1,
        "date": 1,
        "chat": {"id": 10, "type": "private"},
        "from": {"id": 20, "is_bot": False, "first_name": "Test"},
    }
    bot = Bot("123456789:abcdefghijklmnopqrstuvwxyz123456789")
    tg = TelegramAdapter(bot=bot)

    async def native_handlers():
        seen = []

        async def handler(event):
            seen.append(type(event).__name__)

        tg.dispatcher.message.register(handler)
        tg.dispatcher.callback_query.register(handler)
        await tg.dispatcher.feed_update(
            bot,
            Update.model_validate(
                {
                    "update_id": 1,
                    "message": {**payload, "text": "test"},
                }
            ),
        )
        await tg.dispatcher.feed_update(
            bot,
            Update.model_validate(
                {
                    "update_id": 2,
                    "callback_query": {
                        "id": "test",
                        "from": payload["from"],
                        "chat_instance": "1",
                        "inline_message_id": "inline",
                        "data": "test",
                    },
                }
            ),
        )
        require(len(seen) == 2, "native_handlers_intercepted")

    async def media():
        value = {"file_id": "f", "file_unique_id": "u", "width": 10, "height": 10, "duration": 1}
        animation = tg.normalize(Message.model_validate({**payload, "animation": value}))
        note = tg.normalize(Message.model_validate({**payload, "video_note": {**value, "length": 10}}))
        require(bool(animation.attachments) and bool(note.attachments), "animation_or_video_note_missing")

    try:
        for name, operation in (
            ("regression:telegram_native_handlers", native_handlers),
            ("regression:telegram_media", media),
        ):
            report.expect(name)
            await report.check(name, operation)
    finally:
        await tg.aclose()
        await bot.session.close()

    client = SimpleNamespace(request_text=AsyncMock(return_value='{"response":1}'), close=AsyncMock())
    vk = VKAdapter(bot=VKBot(api=API("dummy", http_client=client)))

    async def blocking():
        require(all(not h.blocking for h in vk.labeler.message_view.handlers), "native_handlers_blocked")

    async def callback_ids():
        async def callback(query):
            await vk.edit(query.chat_id, query.message_id, "synthetic test")

        vk.on_callback(callback)
        await vk._receive_callback(
            {
                "object": {
                    "event_id": "event",
                    "user_id": 20,
                    "peer_id": 2000000001,
                    "conversation_message_id": 7,
                    "payload": {"data": "test"},
                }
            }
        )
        params = client.request_text.await_args.kwargs["data"]
        require(
            params.get("cmid") == 7 or params.get("conversation_message_id") == 7,
            "conversation_id_sent_as_global_message_id",
        )

    try:
        for name, operation in (
            ("regression:vk_native_handlers", blocking),
            ("regression:vk_callback_ids", callback_ids),
        ):
            report.expect(name)
            await report.check(name, operation)
    finally:
        await vk.aclose()


def result_details(result):
    require(bool(result.raw.strip()), "empty_llm_response")
    require(result.finish_reason != "length", "llm_response_truncated")
    return {
        "input_tokens": result.input_tokens,
        "output_tokens": result.output_tokens,
        "usage_known": result.usage_known,
        "parse_status": result.parse_status,
        "fallback_used": bool(result.meta.get("fallback_used")),
    }


async def text_probe(llm):
    result = await llm.ainvoke('Return only JSON: {"response":"BOTKIT_OK"}', temperature=0, max_tokens=1024)
    details = result_details(result)
    require(result.response.strip() == "BOTKIT_OK", "unexpected_text_response")
    return details


async def vision_probe(llm, data, mime, tracker):
    service = ContentExtractionService(llm, config=ExtractionConfig(max_vision_pages=1), tracker=tracker)
    result = await service.extract_bytes(data, mime)
    flattened = (result.text or "") + json.dumps(result.structured_data or {})
    require("7319" in flattened and "42" in flattened, "vision_fixture_numbers_not_recognized")
    if mime == "application/pdf":
        require(result.metadata.get("vision_pages") == 1, "scanned_pdf_did_not_use_vision")
    return content_details(result)


async def model_checks(report, chain, fixtures, tracker):
    providers = chain.providers if isinstance(chain, FallbackLLM) else [chain]
    # Test each configured provider directly, so fallback cannot hide a broken Qwen.
    for index, provider in enumerate(providers):
        prefix = f"provider_{index}"
        for suffix, operation in (
            ("text", lambda p=provider: text_probe(p)),
            ("image", lambda p=provider: vision_probe(p, *fixtures["sample.png"], tracker)),
            ("scanned_pdf", lambda p=provider: vision_probe(p, *fixtures["scan.pdf"], tracker)),
        ):
            name = prefix + ":" + suffix
            report.expect(name)
            await report.check(name, operation, timeout=provider.config.timeout + 10)
    report.expect("configured_chain:text", "forced_fallback:text", "forced_fallback:image")
    await report.check(
        "configured_chain:text", lambda: text_probe(chain), timeout=providers[0].config.timeout + 10
    )
    if len(providers) < 2:
        report.record("forced_fallback", "failed", reason="configure_at_least_one_backup_model")
        return

    # Inject HTTP 503 only at the primary's local HTTP boundary. Never disable a real server.
    # The backups are the actual factory-configured live providers, in their actual order.
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(503))) as client:
        primary = LLMGateway(replace(providers[0].config, max_retries=0), client=client, tracker=tracker)
        forced = FallbackLLM(primary, providers[1:], tracker=tracker, timeout=providers[0].config.timeout)

        async def fallback_text():
            details = await text_probe(forced)
            require(details["fallback_used"], "fallback_not_used")
            return {**details, "primary_failure": "injected_http_503", "backup": "live"}

        async def fallback_image():
            before = len(await tracker.events(kind="llm"))
            details = await vision_probe(forced, *fixtures["sample.png"], tracker)
            rows = (await tracker.events(kind="llm"))[before:]
            require(
                any(r["status"] == "success" and r["meta"].get("fallback_used") for r in rows),
                "vision_fallback_not_recorded",
            )
            return {**details, "primary_failure": "injected_http_503", "backup": "live"}

        try:
            await report.check(
                "forced_fallback:text", fallback_text, timeout=providers[0].config.timeout + 10
            )
            await report.check(
                "forced_fallback:image", fallback_image, timeout=providers[0].config.timeout + 10
            )
        finally:
            await primary.aclose()  # the outer owner closes the shared backup providers


async def usage_check(tracker, *, require_llm=False):
    rows = await tracker.events()
    require(bool(rows) and tracker.write_failures == 0, "usage_events_missing_or_write_failed")
    require(all(not r["input_text"] and not r["output_text"] for r in rows), "raw_text_logging_enabled")
    require(all(r["bot_id"] == "acceptance" and r["skill_context"] for r in rows), "usage_context_missing")
    llm = [r for r in rows if r["kind"] == "llm"]
    if require_llm:
        require(any(r["status"] == "success" for r in llm), "no_successful_llm_usage_event")
    return {
        "events": len(rows),
        "llm_events": len(llm),
        "unknown_cost_calls": sum(r["cost"] is None for r in llm),
    }


def successful_llm_check(report):
    """Don't require successful model usage if startup failed before any successful model call."""
    return any(
        case["status"] == "passed"
        and (
            case["name"].startswith(("provider_", "configured_chain:", "forced_fallback:"))
            or case["name"] in {"messenger:llm", "messenger:image_in", "messenger:pdf_scan_in"}
        )
        for case in report.data["cases"]
    )


class BudgetLLM:
    """Limit logical user-triggered calls; retries/fallback can multiply physical calls."""

    def __init__(self, llm, limit):
        self.llm, self.limit, self.used = llm, limit, 0

    async def _call(self, method, *args, **kwargs):
        require(self.used < self.limit, "session_llm_budget_exhausted")
        self.used += 1  # no await before reservation
        return await getattr(self.llm, method)(*args, **kwargs)

    async def ainvoke(self, *args, **kwargs):
        return await self._call("ainvoke", *args, **kwargs)

    async def ainvoke_with_image_b64(self, *args, **kwargs):
        return await self._call("ainvoke_with_image_b64", *args, **kwargs)


class MessengerSession:
    def __init__(self, adapter, llm, tracker, report, fixtures, *, max_calls=10):
        self.adapter, self.tracker, self.report, self.fixtures = adapter, tracker, report, fixtures
        self.llm = BudgetLLM(llm, max_calls)
        self.extraction = ContentExtractionService(
            self.llm,
            tracker=tracker,
            config=replace(ExtractionConfig.from_env(), max_pages=3, max_vision_pages=2),
        )
        self.done, self.lock = asyncio.Event(), asyncio.Lock()
        self.seen = set()
        self.callback_messages = {}  # chat_id -> (nonce, message_id); never share targets across chats
        report.expect(
            "messenger:command",
            "messenger:text",
            "messenger:send",
            "messenger:long_text",
            "messenger:typing",
            "messenger:edit",
            "messenger:delete",
            "messenger:photo_out",
            "messenger:docx_out",
            "messenger:callback",
            "messenger:llm",
            "messenger:pdf_text_in",
            "messenger:pdf_scan_in",
            "messenger:docx_in",
            "messenger:image_in",
            "messenger:finish",
        )
        adapter.on_message(self.receive)
        for name in ("start", "smoke", "llm", "status", "finish"):
            adapter.on_command(name, self.receive)
        adapter.on_callback(self.callback)

    async def tell(self, chat_id, text):
        await self.adapter.send(chat_id, text)

    async def receive(self, message):
        async with self.lock:
            key = (message.chat_id, message.message_id)
            if message.message_id and key in self.seen:
                self.report.record("messenger:duplicate_message", "failed")
                return
            if message.message_id:
                self.seen.add(key)
            await self.report.check("messenger:handle_event", lambda: self._receive(message), timeout=600)

    async def _receive(self, message):
        command = (
            message.text.strip().split(maxsplit=1)[0].split("@", 1)[0].lower() if message.text.strip() else ""
        )
        if command.startswith("/"):
            self.report.record("messenger:command")
        if command == "/start":
            await self.tell(
                message.chat_id,
                "Тест botkit: /smoke → нажмите кнопку → отправьте обычный текст → /llm. "
                "Затем отправьте PDF, DOCX и PNG из папки fixtures (или свои тестовые файлы). "
                "/status — результаты; /finish — сохранить отчёт и остановиться. "
                "Изображения и сканы передаются настроенной модели; не отправляйте секретные файлы.",
            )
        elif command == "/smoke":
            await self.smoke(message.chat_id)
        elif command == "/llm":
            passed = await self.report.check("messenger:llm", lambda: text_probe(self.llm))
            await self.tell(
                message.chat_id, "LLM: PASS" if passed else "LLM: FAIL — см. локальный report.json"
            )
        elif command == "/status":
            await self.tell(message.chat_id, "Ещё не пройдено: " + ", ".join(self.report.data["missing"]))
        elif command == "/finish":
            await self.tell(
                message.chat_id,
                "Тест остановлен. Результат и невыполненные проверки — в локальном report.json.",
            )
            self.report.record("messenger:finish")
            self.done.set()
        elif message.attachments:
            require(len(message.attachments) <= 3, "too_many_attachments_for_smoke_test")
            for attachment in message.attachments:

                async def extract(attachment=attachment):
                    result = await self.extraction.extract(attachment)
                    require(bool((result.text or "").strip()), "empty_attachment_extraction")
                    if kind == "pdf":
                        require(bool(result.metadata.get("page_count")), "pdf_not_recognized_as_pdf")
                        case = "pdf_scan_in" if result.metadata.get("vision_pages") else "pdf_text_in"
                        self.report.record("messenger:" + case, **content_details(result))
                    return content_details(result)

                mime = attachment.mime_type
                suffix = Path(attachment.filename or "").suffix.lower()
                kind = (
                    "pdf"
                    if mime == "application/pdf" or suffix == ".pdf"
                    else (
                        "docx"
                        if mime == DOCX or suffix == ".docx"
                        else "image"
                        if mime.startswith("image/")
                        else "other"
                    )
                )
                passed = await self.report.check("messenger:" + kind + "_in", extract, timeout=400)
                await self.tell(
                    message.chat_id,
                    f"Вложение ({kind}): {'PASS' if passed else 'FAIL'}. Текст файла в отчёт не записывается.",
                )
        elif message.text.strip() and not command.startswith("/"):
            self.report.record("messenger:text")
            await self.tell(message.chat_id, "Текст получен один раз. Для проверки модели: /llm")
        else:
            await self.tell(
                message.chat_id,
                "Не удалось распознать поддерживаемый текст/вложение. /start — шаги проверки.",
            )

    async def smoke(self, chat_id):
        async def send():
            ids = await self.adapter.send(chat_id, "botkit: тестовое сообщение")
            require(bool(ids), "send_returned_no_message_id")

        async def long_text():
            ids = await self.adapter.send(chat_id, "Тест длинного текста 😀\n" * 250)
            require(len(ids) >= 2 and len(ids) == len(set(ids)), "long_message_not_split_or_ids_repeated")

        async def typing():
            # Call the same low-level operation: background typing swallows API failures.
            await self.adapter._operation("typing", lambda: self.adapter._type(chat_id))

        async def edit_delete():
            ids = await self.adapter.send(chat_id, "botkit: это сообщение будет изменено и удалено")
            require(len(ids) == 1 and ids[0].isdigit() and int(ids[0]) > 0, "unexpected_edit_target")
            # Only IDs returned by OUR send; never pass VK callback-local IDs to deletion.
            await self.adapter.edit(chat_id, ids[0], "botkit: сообщение изменено")
            self.report.record("messenger:edit")
            result = await self.adapter.delete(chat_id, ids[0])
            require(
                result is True or (isinstance(result, dict) and result.get(str(ids[0])) == 1),
                "deletion_not_confirmed_by_api",
            )
            self.report.record("messenger:delete")

        for name, operation in (
            ("send", send),
            ("long_text", long_text),
            ("typing", typing),
            ("edit_delete", edit_delete),
        ):
            await self.report.check("messenger:" + name, operation)
        for name, kind in (("sample.png", "photo"), ("sample.docx", "document")):

            async def upload(name=name, kind=kind):
                await self.adapter.send_attachment(chat_id, self.fixtures[name][0], filename=name, kind=kind)

            await self.report.check("messenger:" + ("photo_out" if kind == "photo" else "docx_out"), upload)
        nonce = uuid4().hex[:16]
        ids = await self.adapter.send(
            chat_id, "Нажмите для проверки callback", buttons=[[Button("Проверить", data=nonce)]]
        )
        self.callback_messages[chat_id] = (nonce, ids[-1])

    async def callback(self, query):
        async with self.lock:
            target = self.callback_messages.get(query.chat_id)
            if target is None or query.data != target[0]:
                return

            async def answer():
                await self.adapter.answer_callback(query, "Callback получен")
                # Do not edit/delete using query.message_id: known VK ID namespace defect.
                await self.tell(
                    query.chat_id, "Callback: PASS. Соответствие ID проверяется отдельно в режиме offline."
                )

            await self.report.check("messenger:callback", answer)


async def messenger_checks(args, report, chain, tracker, fixtures):
    from botkit.runtime import make_adapter

    token = os.getenv(args.mode.upper() + "_BOT_TOKEN", "")
    require(bool(token.strip()), "messenger_token_missing")
    adapter = make_adapter(args.mode, token, tracker=tracker, bot_id="acceptance")
    try:
        if args.mode == "telegram":
            info = await adapter.call_api("get_webhook_info")
            require(not info.url, "existing_webhook_stop_or_use_dedicated_test_bot")
        session = MessengerSession(adapter, chain, tracker, report, fixtures, max_calls=args.max_calls)
        report.data["messenger_access"] = "all_users_and_chats"
        report.data["max_logical_model_calls"] = args.max_calls
        report.save()
        print(
            "Polling active: all users/chats accepted. Send /start to the bot. "
            f"Shared model-call limit: {args.max_calls}. /finish stops this test process.",
            flush=True,
        )
        polling = asyncio.create_task(adapter.run())
        finish = asyncio.create_task(session.done.wait())
        try:
            complete, _ = await asyncio.wait(
                {polling, finish}, timeout=args.seconds, return_when=asyncio.FIRST_COMPLETED
            )
            if polling in complete:
                await polling  # propagate unexpected polling failures
                require(session.done.is_set(), "polling_stopped_before_finish")
            elif not complete:
                report.record(
                    "messenger:session_timeout", "failed", reason="finish_not_received_before_deadline"
                )
            else:
                # Let the current callback finish logging its transport event before cancellation.
                async with session.lock:
                    pass
        finally:
            for task in (polling, finish):
                if not task.done():
                    task.cancel()
            await asyncio.gather(polling, finish, return_exceptions=True)
    finally:
        await adapter.aclose()


def parser():
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--mode", choices=["check", "offline", "models", "telegram", "vk"], default="check")
    result.add_argument("--env", type=Path, default=ROOT / ".env.acceptance")
    result.add_argument(
        "--live",
        action="store_true",
        help="Allow paid model requests / replies to anyone contacting this bot",
    )
    result.add_argument(
        "--seconds", type=int, default=900, help="Messenger session lifetime (default 15 minutes)"
    )
    result.add_argument(
        "--max-calls", type=int, default=10, help="Maximum logical LLM calls in messenger mode"
    )
    return result


async def run(args, report):
    if args.mode != "offline":
        require(args.env.is_file(), "env_file_missing_copy_env_acceptance_example")
        load_dotenv(args.env, override=True)
    prices = os.getenv("PRICE_FILE", "config/prices.json")
    price_path = Path(prices)
    if not price_path.is_absolute():
        price_path = ROOT / price_path
    tracker = SQLiteUsageTracker(
        report.directory / "usage.sqlite3", prices=PriceBook(price_path), log_text=False
    )
    chain = None
    try:
        if args.mode != "offline":
            chain = load_llm(tracker=tracker)
            providers = chain.providers if isinstance(chain, FallbackLLM) else [chain]
            report.data["providers"] = [
                {
                    "index": i,
                    "provider": p.config.provider,
                    "model": p.config.model,
                    "streaming_enabled": p.config.supports_streaming,
                    "interval_seconds": p.config.request_interval,
                    "verify_ssl": p.config.verify_ssl,
                    "max_retries": p.config.max_retries,
                }
                for i, p in enumerate(providers)
            ]
            for provider in providers:
                if provider.config.provider == "hub":
                    require(bool(provider.config.api_key.strip()), "hub_api_key_missing")
                    require(
                        provider.config.request_interval >= 2, "hub_interval_must_be_at_least_two_seconds"
                    )
                    require(not provider.config.supports_streaming, "hub_streaming_must_be_disabled")
            report.record("configuration")
        if args.mode == "check":
            # Never contact services, poll, or print tokens/URLs in this mode.
            report.data["live_checks_performed"] = False
            for platform in ("telegram", "vk"):
                report.data[platform + "_token_present"] = bool(os.getenv(platform.upper() + "_BOT_TOKEN"))
            return
        fixtures = make_fixtures(report.directory / "fixtures")
        await local_checks(report, fixtures, tracker)
        if args.mode == "offline":
            await regression_checks(report)
        elif args.mode == "models":
            await model_checks(report, chain, fixtures, tracker)
        else:
            with usage_context(bot_id="acceptance", skill_context="messenger_session"):
                await messenger_checks(args, report, chain, tracker, fixtures)
    finally:
        if chain is not None:
            await chain.aclose()
        if args.mode != "check":
            report.expect("usage:integrity")
            await report.check(
                "usage:integrity", lambda: usage_check(tracker, require_llm=successful_llm_check(report))
            )
        report.data["usage"] = asdict(await tracker.stats())
        report.save()


async def main(argv=None):
    args = parser().parse_args(argv)
    if args.mode in {"models", "telegram", "vk"} and not args.live:
        parser().error("Real requests require --live. They may incur charges and send test messages.")
    if args.seconds <= 0 or args.max_calls <= 0:
        parser().error("--seconds and --max-calls must be positive")
    # Third-party SDK logging may include complete updates, tokens or URLs.
    logging.disable(logging.CRITICAL)
    with contextlib.suppress(ImportError):
        from loguru import logger

        logger.remove()
    directory = (
        ROOT
        / "artifacts"
        / ("acceptance-" + datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8])
    )
    report = Report(directory, args.mode)
    report.expect("run_completed")
    try:
        await report.check("run_completed", lambda: run(args, report), timeout=args.seconds + 1800)
    except asyncio.CancelledError:
        report.record("interrupted", "failed")
        raise
    finally:
        report.save()
        print(f"Report: {directory / 'report.json'}", flush=True)
        print(f"Result: {report.data['result']}; publication_allowed=false", flush=True)
    return report.exit_code


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        raise SystemExit(asyncio.run(main()))
    except KeyboardInterrupt:
        raise SystemExit(130) from None
