from __future__ import annotations

import asyncio
import json
from html.parser import HTMLParser
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from .base import ExtractedContent
from .config import ExtractionConfig, ExtractionError, ExtractionLimitError, UnsupportedContentError

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def inspect_archive(data: bytes, config: ExtractionConfig) -> set[str]:
    """Bound OOXML decompression before any Office parser loads XML/media."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > config.max_archive_entries
                or sum(i.file_size for i in entries) > config.max_archive_bytes
            ):
                raise ExtractionLimitError("Office archive exceeds decompression limits")
            if len({i.filename for i in entries}) != len(entries):
                raise ExtractionError("Duplicate Office archive members")
            for info in entries:
                if info.flag_bits & 1:
                    raise ExtractionError("Encrypted Office archives are not supported")
                if "vbaproject" in info.filename.lower():
                    raise UnsupportedContentError("Macro-enabled documents are not supported")
                if info.filename.endswith((".xml", ".rels")):
                    xml = archive.read(info)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise ExtractionError("XML entities are not allowed in Office documents")
            return {i.filename for i in entries}
    except BadZipFile:
        raise ExtractionError("Invalid Office archive") from None


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "br", "div", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


class TextExtractor:
    kind = "document"
    mime_types = {
        "text/plain",
        "text/markdown",
        "text/csv",
        "text/tab-separated-values",
        "text/html",
        "application/json",
        "text/json",
        "text/xml",
        "application/xml",
    }

    def __init__(self, config: ExtractionConfig | None = None):
        self.config = config or ExtractionConfig()

    async def supports(self, mime_type: str) -> bool:
        return mime_type in self.mime_types

    def _extract(self, data: bytes, mime_type: str) -> ExtractedContent:
        self.config.check_bytes(data)
        warnings = []
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        if data.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
            encoding = "utf-32"
        try:
            text = data.decode(encoding)
        except UnicodeError:
            try:
                text, encoding = data.decode("cp1251"), "cp1251"
                warnings.append("encoding_fallback_cp1251_verify")
            except UnicodeError:
                raise ExtractionError("Unsupported text encoding; upload UTF-8") from None
        if "\x00" in text or any(ord(c) < 32 and c not in "\t\n\r\f" for c in text):
            raise ExtractionError("Binary data cannot be extracted as text")
        self.config.check_text(text)
        structured = None
        if mime_type in {"application/json", "text/json"}:
            try:
                structured = {"value": json.loads(text)}
            except (ValueError, RecursionError):
                warnings.append("invalid_json_preserved_as_text")
        elif mime_type == "text/html":
            parser = _HTMLText()
            parser.feed(text)
            text = "".join(parser.parts).strip()
            warnings.append("html_text_only_no_external_resources")
        if not text.strip():
            warnings.append("empty_extraction")
        return ExtractedContent(
            "structured" if structured is not None else "text",
            text,
            structured,
            warnings,
            "low" if warnings else "high",
            {"encoding": encoding, "method": "native"},
        )

    async def extract(self, data: bytes, mime_type: str) -> ExtractedContent:
        return await asyncio.to_thread(self._extract, data, mime_type)


class OfficeExtractor:
    kind = "document"

    def __init__(self, config: ExtractionConfig | None = None):
        self.config = config or ExtractionConfig()

    async def supports(self, mime_type: str) -> bool:
        return mime_type in {DOCX, XLSX, PPTX}

    def _docx(self, data: bytes) -> tuple[str, dict, list[str]]:
        from docx import Document
        from docx.table import Table

        document = Document(BytesIO(data))
        blocks, tables = [], []
        cells = 0

        def walk(container, depth=0):
            nonlocal cells
            if depth > 12:
                raise ExtractionLimitError("DOCX nesting limit exceeded")
            for item in container.iter_inner_content():
                if isinstance(item, Table):
                    rows = []
                    for row in item.rows:
                        cells += len(row.cells)
                        if cells > self.config.max_cells:
                            raise ExtractionLimitError("DOCX cell limit exceeded")
                        rows.append(["\n".join(walk(cell, depth + 1)) for cell in row.cells])
                    tables.append({"rows": rows})
                    yield "\n".join("\t".join(row) for row in rows)
                else:
                    yield item.text

        blocks.extend(walk(document))
        return "\n".join(blocks), {"tables": tables}, ["docx_body_only_headers_footers_notes_not_extracted"]

    def _xlsx(self, data: bytes) -> tuple[str, dict, list[str]]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UnsupportedContentError('XLSX requires pip install "botkit[office]"') from None
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheets, cells = [], 0
        try:
            if len(workbook.worksheets) > self.config.max_pages:
                raise ExtractionLimitError("Workbook sheet limit exceeded")
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > self.config.max_cells:
                    raise ExtractionLimitError("Worksheet dimensions exceed cell limit")
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > self.config.max_cells:
                        raise ExtractionLimitError("Workbook cell limit exceeded")
                    rows.append(["" if value is None else str(value) for value in row])
                sheets.append({"name": sheet.title, "rows": rows})
        finally:
            workbook.close()
        text = "\n\n".join(
            f"[Sheet: {s['name']}]\n" + "\n".join("\t".join(r) for r in s["rows"]) for s in sheets
        )
        return text, {"sheets": sheets}, ["xlsx_formulas_not_evaluated", "office_visuals_not_extracted"]

    def _pptx(self, data: bytes) -> tuple[str, dict, list[str]]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError:
            raise UnsupportedContentError('PPTX requires pip install "botkit[office]"') from None
        presentation = Presentation(BytesIO(data))
        if len(presentation.slides) > self.config.max_pages:
            raise ExtractionLimitError("Presentation slide limit exceeded")
        slides, cells = [], 0

        def texts(shapes, depth=0):
            nonlocal cells
            if depth > 12:
                raise ExtractionLimitError("Slide nesting limit exceeded")
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from texts(shape.shapes, depth + 1)
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells += len(row.cells)
                        if cells > self.config.max_cells:
                            raise ExtractionLimitError("Presentation cell limit exceeded")
                        yield "\t".join(c.text for c in row.cells)
                elif shape.has_text_frame:
                    yield shape.text_frame.text

        for index, slide in enumerate(presentation.slides, 1):
            slides.append({"slide_number": index, "text": "\n".join(texts(slide.shapes))})
        return (
            "\n\n".join(f"[Slide {s['slide_number']}]\n{s['text']}" for s in slides),
            {"slides": slides},
            ["pptx_shape_order_not_visual_order", "office_visuals_and_speaker_notes_not_extracted"],
        )

    def _extract(self, data: bytes, mime_type: str) -> ExtractedContent:
        self.config.check_bytes(data)
        names = inspect_archive(data, self.config)
        required = {DOCX: "word/document.xml", XLSX: "xl/workbook.xml", PPTX: "ppt/presentation.xml"}
        if required.get(mime_type) not in names:
            raise ExtractionError("Office content does not match document type")
        try:
            text, structured, warnings = {DOCX: self._docx, XLSX: self._xlsx, PPTX: self._pptx}[mime_type](
                data
            )
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError("Office document could not be parsed") from None
        if mime_type == DOCX and any(n.startswith("word/media/") for n in names):
            warnings.append("embedded_images_not_extracted_export_pdf_for_vision")
        self.config.check_text(text)
        if not text.strip():
            warnings.append("empty_extraction")
        return ExtractedContent(
            "text", text, structured, warnings, "low" if warnings else "high", {"method": "native"}
        )

    async def extract(self, data: bytes, mime_type: str) -> ExtractedContent:
        return await asyncio.to_thread(self._extract, data, mime_type)
